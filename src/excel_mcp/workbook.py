import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .exceptions import WorkbookError

logger = logging.getLogger(__name__)

DEFAULT_SHEET_NAME = "Sheet1"
# Appended to a tool's success message when add_sheet took over the default sheet,
# so the model does not look for "Sheet1" afterwards.
REPLACED_DEFAULT_SHEET_NOTE = f" (it replaces the empty default sheet {DEFAULT_SHEET_NAME})"

def create_workbook(filepath: str, sheet_name: str = DEFAULT_SHEET_NAME) -> dict[str, Any]:
    """Create a new Excel workbook with optional custom sheet name"""
    try:
        wb = Workbook()
        # Rename default sheet
        if "Sheet" in wb.sheetnames:
            sheet = wb["Sheet"]
            sheet.title = sheet_name
        else:
            wb.create_sheet(sheet_name)

        path = Path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return {
            "message": f"Created workbook: {filepath}",
            "active_sheet": sheet_name,
            "workbook": wb
        }
    except Exception as e:
        logger.error(f"Failed to create workbook: {e}")
        raise WorkbookError(f"Failed to create workbook: {e!s}")

def get_or_create_workbook(filepath: str) -> Workbook:
    """Get existing workbook or create new one if it doesn't exist"""
    try:
        return load_workbook(filepath)
    except FileNotFoundError:
        return create_workbook(filepath)["workbook"]

def _is_untouched_default_sheet(wb: Workbook) -> bool:
    """The workbook's only sheet is the default "Sheet1" with nothing put on it.

    Renaming the sheet would carry whatever it holds over to the new sheet, so any
    cell (a value, a style, a comment or a hyperlink), merged range, chart, image,
    table, data validation, conditional format, print area or print titles, frozen
    panes, auto filter, sheet protection or tab colour counts as work. So does a
    defined name scoped to the sheet or referring to it, because openpyxl does not
    update defined names when a sheet is renamed.
    """
    if wb.sheetnames != [DEFAULT_SHEET_NAME]:
        return False
    sheet = wb[DEFAULT_SHEET_NAME]
    return (
        isinstance(sheet, Worksheet)
        and not sheet._cells
        and not sheet.merged_cells
        and not sheet._charts
        and not sheet._images
        and not sheet.tables
        and not sheet.data_validations.dataValidation
        and not sheet.conditional_formatting
        and not sheet.print_area
        and not sheet.print_titles
        and not sheet.freeze_panes
        and not sheet.auto_filter.ref
        and not sheet.protection.sheet
        and not sheet.sheet_properties.tabColor
        and not sheet.defined_names
        and not any(
            DEFAULT_SHEET_NAME.lower() in (name.value or "").lower()
            for name in wb.defined_names.values()
        )
    )

def add_sheet(wb: Workbook, sheet_name: str) -> bool:
    """Add a worksheet named ``sheet_name``; True when it took over the default sheet.

    A new workbook starts with an empty default "Sheet1". While that is still its
    only sheet and untouched, it is renamed to ``sheet_name`` instead of staying
    in front of the new sheet as an empty first tab.
    """
    if _is_untouched_default_sheet(wb):
        wb[DEFAULT_SHEET_NAME].title = sheet_name
        return True
    wb.create_sheet(sheet_name)
    return False

def create_sheet(filepath: str, sheet_name: str) -> dict:
    """Create a new worksheet in the workbook if it doesn't exist."""
    try:
        wb = load_workbook(filepath)

        # Check if sheet already exists
        if sheet_name in wb.sheetnames:
            raise WorkbookError(f"Sheet {sheet_name} already exists")

        # Create new sheet, or take over the untouched default one
        replaced_default = add_sheet(wb, sheet_name)
        wb.save(filepath)
        wb.close()
        message = f"Sheet {sheet_name} created successfully"
        if replaced_default:
            message += REPLACED_DEFAULT_SHEET_NOTE
        return {"message": message}
    except WorkbookError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to create sheet: {e}")
        raise WorkbookError(str(e))

def get_workbook_info(filepath: str, include_ranges: bool = False) -> dict[str, Any]:
    """Get metadata about workbook including sheets, ranges, etc."""
    try:
        path = Path(filepath)
        if not path.exists():
            raise WorkbookError(f"File not found: {filepath}")
            
        wb = load_workbook(filepath, read_only=False)
        
        info = {
            "filename": path.name,
            "sheets": wb.sheetnames,
            "size": path.stat().st_size,
            "modified": path.stat().st_mtime
        }
        
        if include_ranges:
            # Add used ranges for each sheet
            ranges = {}
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if ws.max_row > 0 and ws.max_column > 0:
                    ranges[sheet_name] = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            info["used_ranges"] = ranges
            
        wb.close()
        return info
        
    except WorkbookError as e:
        logger.error(str(e))
        raise
    except Exception as e:
        logger.error(f"Failed to get workbook info: {e}")
        raise WorkbookError(str(e))
