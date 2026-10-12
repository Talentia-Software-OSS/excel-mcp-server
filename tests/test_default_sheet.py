import os
import sys
import tempfile
import unittest

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.chart import BarChart  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.table import Table  # noqa: E402

import excel_mcp.server as server  # noqa: E402
from excel_mcp.data import write_data  # noqa: E402
from excel_mcp.workbook import add_sheet, create_sheet, create_workbook  # noqa: E402


class TestDefaultSheet(unittest.TestCase):
    """An untouched default "Sheet1" is taken over by the first named sheet
    instead of staying in the workbook as an empty first tab."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self._tmp.name, "book.xlsx")
        create_workbook(self.path)

    def tearDown(self):
        self._tmp.cleanup()

    def _sheetnames(self):
        wb = load_workbook(self.path)
        names = wb.sheetnames
        wb.close()
        return names

    def _change_sheet1(self, change):
        wb = load_workbook(self.path)
        change(wb["Sheet1"])
        wb.save(self.path)
        wb.close()

    def test_create_sheet_renames_the_untouched_default_sheet(self):
        result = create_sheet(self.path, "Revenue")

        self.assertEqual(self._sheetnames(), ["Revenue"])
        self.assertIn("replaces the empty default sheet Sheet1", result["message"])

    def test_write_data_renames_the_untouched_default_sheet(self):
        write_data(self.path, "Revenue", [["Region", "Revenue"], ["North", 120]])

        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["Revenue"])
        self.assertEqual(wb["Revenue"]["B2"].value, 120)
        wb.close()

    def test_a_default_sheet_holding_work_is_kept(self):
        changes = {
            "value": lambda ws: ws.cell(row=1, column=1, value="kept"),
            "merged cells": lambda ws: ws.merge_cells("A1:B2"),
            "data validation": lambda ws: ws.add_data_validation(
                DataValidation(type="list", formula1='"Yes,No"', sqref="A1"),
            ),
            "conditional format": lambda ws: ws.conditional_formatting.add(
                "A1:A5",
                CellIsRule(
                    operator="greaterThan", formula=["0"],
                    fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
                ),
            ),
            "style only": lambda ws: setattr(ws["A1"], "font", Font(bold=True)),
            "comment": lambda ws: setattr(ws["A1"], "comment", Comment("Check this", "Analyst")),
            "defined name": lambda ws: ws.parent.defined_names.add(
                DefinedName("InputCell", attr_text="Sheet1!$A$1"),
            ),
            "sheet-scoped name": lambda ws: ws.defined_names.add(
                DefinedName("TaxRate", attr_text="0.2"),
            ),
            "print area": lambda ws: setattr(ws, "print_area", "A1:C10"),
            "print titles": lambda ws: setattr(ws, "print_title_rows", "1:1"),
            "frozen panes": lambda ws: setattr(ws, "freeze_panes", "A2"),
            "auto filter": lambda ws: setattr(ws.auto_filter, "ref", "A1:C1"),
            "sheet protection": lambda ws: ws.protection.enable(),
            "tab colour": lambda ws: setattr(ws.sheet_properties, "tabColor", "FF0000"),
        }
        for kind, change in changes.items():
            with self.subTest(kind):
                create_workbook(self.path)
                self._change_sheet1(change)

                result = create_sheet(self.path, "Revenue")

                self.assertEqual(self._sheetnames(), ["Sheet1", "Revenue"])
                self.assertEqual(result["message"], "Sheet Revenue created successfully")

    def test_a_default_sheet_with_a_chart_or_a_table_is_kept(self):
        changes = {
            "chart": lambda ws: ws.add_chart(BarChart(), "C3"),
            "table": lambda ws: ws.add_table(Table(displayName="Regions", ref="A1:B2")),
        }
        for kind, change in changes.items():
            with self.subTest(kind):
                wb = Workbook()
                wb.active.title = "Sheet1"
                change(wb.active)

                self.assertFalse(add_sheet(wb, "Revenue"))
                self.assertEqual(wb.sheetnames, ["Sheet1", "Revenue"])

    def test_a_workbook_with_two_sheets_keeps_both(self):
        wb = load_workbook(self.path)
        wb.create_sheet("Notes")
        wb.save(self.path)
        wb.close()

        create_sheet(self.path, "Revenue")
        write_data(self.path, "Costs", [["Region", "Cost"]])

        self.assertEqual(self._sheetnames(), ["Sheet1", "Notes", "Revenue", "Costs"])

    def test_an_untouched_first_sheet_with_its_own_name_is_kept(self):
        create_workbook(self.path, sheet_name="Data")

        create_sheet(self.path, "Revenue")

        self.assertEqual(self._sheetnames(), ["Data", "Revenue"])


class TestCreateWorkbookTool(unittest.TestCase):
    def setUp(self):
        server.EXCEL_FILES_PATH = None
        self._tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self._tmp.name, "book.xlsx")

    def tearDown(self):
        self._tmp.cleanup()

    def test_names_the_first_sheet(self):
        server.create_workbook(self.path, sheet_name="Revenue")

        self.assertEqual(load_workbook(self.path).sheetnames, ["Revenue"])

    def test_the_first_sheet_defaults_to_sheet1(self):
        server.create_workbook(self.path)

        self.assertEqual(load_workbook(self.path).sheetnames, ["Sheet1"])


if __name__ == "__main__":
    unittest.main()
